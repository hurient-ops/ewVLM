import os
import docx

screens = [
    ("Login", "시스템 접근을 위한 사용자 인증 화면입니다. ID와 비밀번호를 입력하여 로그인합니다."),
    ("Signup", "신규 사용자의 계정 생성을 위한 회원가입 화면입니다. 관리자 승인 후 사용 가능합니다."),
    ("MonitorALiveControl", "관제 A 구역: CCTV 실시간 라이브 스트리밍 영상을 다채널로 모니터링하는 메인 화면입니다."),
    ("PtzPatrolSchedule", "PTZ 카메라의 순찰 경로(프리셋) 및 시간대별 스케줄을 설정하고 관리하는 화면입니다."),
    ("PtzTargetHandover", "한 카메라의 시야에서 벗어난 타겟을 인접 PTZ 카메라로 자동 인계하여 연속 추적하는 기능을 설정합니다."),
    ("MobilePatrolApp", "현장 순찰 요원의 모바일 기기와 연동하여 현장 영상 송출 및 양방향 통신을 지원하는 화면입니다."),
    ("IpAudioBroadcastConsole", "IP 스피커를 통해 특정 구역에 자동 경고 방송을 송출하거나 수동 마이크 방송을 제어하는 콘솔입니다."),
    ("HardwareSelfHealingShell", "장애가 발생한 카메라나 에지 디바이스를 원격으로 자동 재부팅하거나 초기화하는 자가 치유 제어 화면입니다."),
    ("MassDeviceConfigClone", "하나의 카메라에 설정된 구성값을 동일 기종의 다수 카메라에 일괄 복제하여 적용하는 화면입니다."),
    ("EdgeAiOrchestration", "에지 디바이스에 경량화된 AI 모델을 원격으로 배포, 업데이트 및 버전 관리하는 오케스트레이션 화면입니다."),
    ("PrivacyExportWorkshop", "외부 반출용 영상에서 사람의 얼굴이나 차량 번호판을 자동/수동으로 모자이크 처리하여 반출하는 워크샵 화면입니다."),
    ("AlertCenterDashboard", "발생한 모든 보안/안전 알람 내역을 타임라인 및 중요도 순으로 통합 관리하는 대시보드입니다."),
    ("MonitorBVlmAnalysis", "관제 B 구역: VLM(시각언어모델)이 분석한 영상의 의미론적 해석 결과(캡션, 위협 수준, 권장 조치)를 실시간 모니터링합니다."),
    ("EventReviewCenter", "과거에 발생한 이벤트 영상을 VLM의 상세 리포트와 함께 다시 보며 사후 분석 및 평가를 진행하는 센터입니다."),
    ("DisasterVirtualWarRoom", "대형 재난 발생 시 관련 부서가 원격으로 동시 접속하여 실시간 영상과 현장 정보를 공유하는 가상 합동 워룸입니다."),
    ("RealtimeBiDashboard", "전체 시스템의 운영 현황, 알람 통계, AI 분석 추이를 시각적 그래프로 제공하는 실시간 BI 대시보드입니다."),
    ("VssSemanticSearch", "키워드나 자연어 문장을 입력하여 녹화된 영상 중 해당 의미와 일치하는 장면을 찾아내는 시맨틱 검색 화면입니다."),
    ("NaturalLanguageRuleCopilot", "사용자가 자연어로 '안전모 안 쓴 사람 찾아줘'라고 입력하면 VLM이 이를 이해하고 알람 규칙으로 자동 생성해 주는 코파일럿 화면입니다."),
    ("SemanticVectorPortal", "AI가 추출한 영상의 특징 벡터 데이터베이스를 관리하고 검색 임계값 등을 세부 튜닝하는 벡터 포털입니다."),
    ("PromptGatewayDeploy", "VLM 모델에 주입되는 시스템 프롬프트를 시나리오별로 작성, 테스트 및 운영 환경에 배포하는 관리 화면입니다."),
    ("LoraFinetuningConsole", "특정 현장에 맞게 VLM의 인식률을 높이기 위해 추가 데이터로 LoRA 파인튜닝을 진행하고 모델을 학습시키는 콘솔입니다."),
    ("GisSmartMap", "전자 지도(GIS) 상에 카메라 위치와 이벤트 발생 현황을 아이콘 및 히트맵 형태로 직관적으로 표시하는 통합 관제 화면입니다."),
    ("MultiChannelSyncPlayback", "다수의 카메라 녹화 영상을 동일한 시간대(타임스탬프)로 동기화하여 동시에 재생하고 추적하는 화면입니다."),
    ("CameraSecurityPortal", "카메라 접속 비밀번호 주기적 변경, 펌웨어 무결성 검증, RTSP 스트림 암호화(SRTP) 등을 설정하는 보안 포털입니다."),
    ("CameraSetupConfig", "개별 카메라의 해상도, 프레임 레이트, 비트레이트, OSD(On-Screen Display) 등 세부 설정을 변경하는 화면입니다."),
    ("CameraListManager", "등록된 전체 카메라의 목록, IP 주소, 연결 상태, 모델명 등을 엑셀 형태의 그리드로 조회하고 관리하는 화면입니다."),
    ("GeometryCalibrationConsole", "카메라의 렌즈 왜곡을 보정하고 2D 영상 좌표를 3D 실제 공간 좌표로 매핑하기 위한 기하학 캘리브레이션 툴입니다."),
    ("NetworkTopologyMonitor", "카메라, 스위치, NVR, 서버 간의 네트워크 연결 상태와 대역폭 사용량을 시각적 토폴로지 맵으로 모니터링합니다."),
    ("NvrStorageDashboard", "NVR 및 스토리지의 디스크 사용량, 녹화 보존 기간, RAID 상태 및 디스크 수명(S.M.A.R.T)을 관리하는 대시보드입니다."),
    ("MultiSiteAuthMatrix", "다수의 현장(사이트)을 통합 관리할 때, 관리자 및 사용자별로 접근 가능한 카메라와 메뉴 권한을 세밀하게 매트릭스 형태로 설정하는 화면입니다."),
    ("SystemAuditLogPortal", "사용자의 로그인 이력, 설정 변경 내역, 영상 조회 기록 등 시스템 내 모든 활동을 기록하고 감사(Audit)하는 포털입니다.")
]

# Create Markdown
md_content = "# ewVLM 전체 시스템 사용 설명서\n\n이 문서에서는 ewVLM 플랫폼의 모든 화면에 대한 설명 및 사용 방법을 제공합니다.\n\n"

for title, desc in screens:
    md_content += f"## {title}\n"
    md_content += f"{desc}\n\n"

with open("e:/projects/ewVLM/ewVLM_User_Manual.md", "w", encoding="utf-8") as f:
    f.write(md_content)

# Create DOCX
doc = docx.Document()
doc.add_heading("ewVLM 전체 시스템 사용 설명서", 0)
doc.add_paragraph("이 문서에서는 ewVLM 플랫폼의 모든 화면에 대한 설명 및 사용 방법을 제공합니다.")

for title, desc in screens:
    doc.add_heading(title, level=2)
    doc.add_paragraph(desc)

doc.save("e:/projects/ewVLM/ewVLM_User_Manual.docx")
print("User manuals generated successfully: ewVLM_User_Manual.md and ewVLM_User_Manual.docx")
